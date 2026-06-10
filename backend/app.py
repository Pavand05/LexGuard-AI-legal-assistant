import sys
import os

# Ensure UTF-8 output on Windows (avoids UnicodeEncodeError for emoji)
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from flask import Flask, request, jsonify
from flask_cors import CORS
# pyrefly: ignore [missing-import]
from flask_jwt_extended import jwt_required, get_jwt_identity
import fitz  # PyMuPDF
from docx import Document as DocxDocument

from extensions import db, jwt, bcrypt
from models import User, Document
from auth import auth_bp
from cases_api import cases_bp
from ai_processor import document_processor

# ──────────────────────────────────────────────────────────────────────────────
# App factory
# ──────────────────────────────────────────────────────────────────────────────
app = Flask(__name__)

# ----- Config ----------------------------------------------------------------
BASE_DIR      = os.path.dirname(os.path.abspath(__file__))
if os.environ.get('VERCEL') or os.environ.get('DATABASE_URL'):
    UPLOAD_FOLDER = '/tmp/uploads'
else:
    UPLOAD_FOLDER = os.path.join(BASE_DIR, 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

db_url = os.environ.get('DATABASE_URL')
if db_url and db_url.startswith("postgres://"):
    db_url = db_url.replace("postgres://", "postgresql://", 1)

if not db_url:
    if os.environ.get('VERCEL'):
        db_url = "sqlite:////tmp/legal_ai.db"
    else:
        db_url = f"sqlite:///{os.path.join(BASE_DIR, 'legal_ai.db')}"

app.config['SQLALCHEMY_DATABASE_URI']        = db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['JWT_SECRET_KEY']                 = os.environ.get('JWT_SECRET_KEY', 'change-me-in-production-32chars!!')
app.config['JWT_ACCESS_TOKEN_EXPIRES']       = 86400   # 24 hours (seconds)
app.config['UPLOAD_FOLDER']                  = UPLOAD_FOLDER
app.config['JWT_TOKEN_LOCATION']             = ['headers', 'query_string']
app.config['JWT_QUERY_STRING_NAME']           = 'token'

# ----- Extensions ------------------------------------------------------------
db.init_app(app)
jwt.init_app(app)
bcrypt.init_app(app)
CORS(app)

# ----- Blueprints ------------------------------------------------------------
app.register_blueprint(auth_bp)
app.register_blueprint(cases_bp)

# ──────────────────────────────────────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────────────────────────────────────
def extract_text(file_path):
    """
    Returns (full_text: str, page_texts: list[str]).
    page_texts is a list of per-page strings (PDFs only); for other formats it is [full_text].
    """
    if file_path.endswith('.pdf'):
        page_texts = []
        with fitz.open(file_path) as doc:
            for page in doc:
                try:
                    page_texts.append(page.get_text("text"))  # type: ignore
                except Exception:
                    page_texts.append(str(page))
        return "".join(page_texts), page_texts

    elif file_path.endswith('.docx'):
        doc  = DocxDocument(file_path)
        text = "\n".join(p.text for p in doc.paragraphs)
        return text, [text]

    elif file_path.endswith('.txt'):
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        return text, [text]

    else:
        return "Unsupported file format", ["Unsupported file format"]


def current_user():
    """Return the User ORM object for the authenticated request."""
    user_id = int(get_jwt_identity())
    return db.session.get(User, user_id)


# ── Global Request Status Guard ───────────────────────────────────────────────
from flask_jwt_extended import verify_jwt_in_request, get_jwt_identity

@app.before_request
def check_user_status():
    # Only enforce if request path is not one of the public endpoints
    # Allowed endpoints for unverified/general users:
    # 1. /health
    # 2. Auth routes like login, register, and register/lawyer, and getting own details /auth/me
    allowed_paths = [
        '/health',
        '/auth/login',
        '/auth/register',
        '/auth/register/lawyer',
        '/auth/me'
    ]
    if request.path in allowed_paths:
        return
        
    try:
        verify_jwt_in_request(optional=True)
        identity = get_jwt_identity()
        if identity:
            user = db.session.get(User, int(identity))
            if user and user.status != 'active':
                return jsonify({"error": "Access denied. Account not verified."}), 403
    except Exception:
        pass


# ──────────────────────────────────────────────────────────────────────────────
# Routes
# ──────────────────────────────────────────────────────────────────────────────

@app.route('/health', methods=['GET'])
def health_check():
    """Health check endpoint (public)."""
    return jsonify({"status": "healthy", "ai_processor": "initialized"})


# ── Upload & process document ─────────────────────────────────────────────────
@app.route('/upload', methods=['POST'])
@jwt_required()
def upload_and_extract():
    user = current_user()
    if not user:
        return jsonify({"error": "User not found"}), 404

    print("Received upload request.")
    file = request.files.get('file')
    if not file:
        return jsonify({"error": "No file provided"}), 400

    filename = file.filename
    if not filename:
        return jsonify({"error": "Invalid filename"}), 400

    print(f"Uploaded file: {filename}")

    # Save to uploads/<user_id>/ folder for isolation
    user_upload_dir = os.path.join(UPLOAD_FOLDER, str(user.id))
    os.makedirs(user_upload_dir, exist_ok=True)
    file_path = os.path.join(user_upload_dir, filename)
    file.save(file_path)
    print(f"Saved to: {file_path}")

    try:
        # Extract text
        extracted_text, page_texts = extract_text(file_path)
        print(f"Text extraction done. Pages: {len(page_texts)}, Length: {len(extracted_text)}")

        # AI processing
        print("Starting AI processing...")
        summary     = document_processor.summarize_document(extracted_text)
        clauses     = document_processor.extract_clauses(extracted_text, page_texts=page_texts)
        risk_stats  = document_processor.analyze_risks(clauses)
        print("AI processing complete")

        # Persist to DB
        doc = Document(
            user_id       = user.id,
            filename      = os.path.join(str(user.id), filename),
            original_name = filename,
            summary       = summary,
            text          = extracted_text,
        )
        doc.set_clauses(clauses)
        doc.set_risks(risk_stats)
        db.session.add(doc)
        db.session.commit()

        analysis_results = {
            "id":          doc.id,
            "summary":     summary,
            "clauses":     clauses,
            "risks":       risk_stats,
            "processed":   True,
            "filename":    filename,
            "text_length": len(extracted_text),
            "text":        extracted_text
        }

        return jsonify({
            "message":  "File uploaded and processed successfully",
            "analysis": analysis_results
        })

    except Exception as e:
        print(f"ERROR: {e}")
        return jsonify({"error": str(e)}), 500


# ── Chat with document ────────────────────────────────────────────────────────
@app.route('/chat', methods=['POST'])
@jwt_required()
def chat_with_document():
    """Handle chat queries about an uploaded document."""
    try:
        data = request.get_json()
        if not data:
            return jsonify({"error": "No JSON data provided"}), 400

        question      = data.get('question', '')
        document_text = data.get('document_text', '')
        use_chroma    = data.get('use_chroma', False)

        if not question or not document_text:
            return jsonify({"error": "Question and document text are required"}), 400

        rag_result = document_processor.rag_qa(question, document_text, use_chroma=use_chroma)

        return jsonify({
            "response":  rag_result["answer"],
            "sources":   rag_result["sources"],
            "question":  question
        })

    except Exception as e:
        print(f"Chat error: {e}")
        return jsonify({"error": str(e)}), 500


# ── Document library ──────────────────────────────────────────────────────────
@app.route('/documents', methods=['GET'])
@jwt_required()
def list_documents():
    """Return all documents belonging to the current user (newest first)."""
    user = current_user()
    if not user:
        return jsonify({"error": "User not found"}), 404

    docs = user.documents.order_by(Document.uploaded_at.desc()).all()
    return jsonify({"documents": [d.to_dict() for d in docs]})


@app.route('/documents/<int:doc_id>', methods=['GET'])
@jwt_required()
def get_document(doc_id):
    """Return a single document (with full text) belonging to the current user."""
    user = current_user()
    if not user:
        return jsonify({"error": "User not found"}), 404

    doc = Document.query.filter_by(id=doc_id, user_id=user.id).first()
    if not doc:
        return jsonify({"error": "Document not found"}), 404

    return jsonify({"document": doc.to_dict(include_text=True)})


@app.route('/documents/<int:doc_id>', methods=['DELETE'])
@jwt_required()
def delete_document(doc_id):
    """Delete a document belonging to the current user."""
    user = current_user()
    if not user:
        return jsonify({"error": "User not found"}), 404

    doc = Document.query.filter_by(id=doc_id, user_id=user.id).first()
    if not doc:
        return jsonify({"error": "Document not found"}), 404

    # Remove physical file if it still exists
    file_path = os.path.join(UPLOAD_FOLDER, doc.filename)
    try:
        if os.path.exists(file_path):
            os.remove(file_path)
    except OSError:
        pass

    db.session.delete(doc)
    db.session.commit()
    return jsonify({"message": "Document deleted successfully"})


# ──────────────────────────────────────────────────────────────────────────────
# Entrypoint
# ──────────────────────────────────────────────────────────────────────────────
with app.app_context():
    try:
        db.create_all()
        print("Database tables created / verified.")
        
        # Seed default admin if not exists
        admin = User.query.filter_by(email="admin@legalai.com").first()
        if not admin:
            pw_hash = bcrypt.generate_password_hash("admin123").decode('utf-8')
            admin = User(name="System Admin", email="admin@legalai.com", password_hash=pw_hash, role="admin", status="active")
            db.session.add(admin)
            db.session.commit()
            print("Default admin account created: admin@legalai.com / admin123")
    except Exception as e:
        print(f"Database/admin initialization error: {e}")

print("Initializing AI Document Processor...")
try:
    document_processor.initialize_llm()
    print("AI Processor initialized successfully")
except Exception as e:
    print(f"AI Processor initialization error: {e}")

if __name__ == '__main__':
    # Bind to host 0.0.0.0 and dynamic environment PORT for cloud environments like Render
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
